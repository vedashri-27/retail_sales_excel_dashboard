"""Copy LibreOffice-computed values into the openpyxl workbook as cached formula results,
so previewers and viewers that don't recalculate still show numbers. Formulas are untouched."""
import sys, zipfile, re
from lxml import etree
from openpyxl import load_workbook

src, calc, out = sys.argv[1:4]
vals = load_workbook(calc, data_only=True)
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
sheet_files = {f"xl/worksheets/sheet{i+1}.xml": name for i, name in enumerate(vals.sheetnames)}

zin = zipfile.ZipFile(src)
zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
filled = 0
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename in sheet_files:
        ws = vals[sheet_files[item.filename]]
        root = etree.fromstring(data)
        for c in root.iter(f"{{{NS}}}c"):
            fe = c.find(f"{{{NS}}}f")
            if fe is None:
                continue
            v = ws[c.get("r")].value
            if v is None:
                v = ""
            ve = c.find(f"{{{NS}}}v")
            if ve is None:
                ve = etree.SubElement(c, f"{{{NS}}}v")
            if isinstance(v, bool):
                c.set("t", "b"); ve.text = "1" if v else "0"
            elif isinstance(v, (int, float)):
                c.attrib.pop("t", None); ve.text = repr(float(v)) if isinstance(v, float) else str(v)
            else:
                c.set("t", "str"); ve.text = str(v)
            filled += 1
        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    zout.writestr(item, data)
zout.close()
print("cached values written:", filled)
